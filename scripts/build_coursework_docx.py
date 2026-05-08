from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = BASE_DIR / "COURSEWORK_REPORT.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(hdr[index], header, bold=True)
        set_cell_shading(hdr[index], "D9EAF7")
        hdr[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Normal"].paragraph_format.space_after = Pt(6)

    for style_name, size, color in (
        ("Title", 20, RGBColor(31, 78, 121)),
        ("Heading 1", 15, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def build() -> None:
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("КУРСОВАЯ РАБОТА")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Разработка локальной OSINT-системы для сбора и анализа открытых данных пользователей Telegram"
    )
    run.font.size = Pt(14)

    document.add_paragraph()
    meta_rows = [
        ("Студент", "____________________________"),
        ("Группа", "____________________________"),
        ("Преподаватель", "____________________________"),
        ("Год", "2026"),
    ]
    meta = document.add_table(rows=1, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    for row_index, (left, right) in enumerate(meta_rows):
        cells = meta.rows[0].cells if row_index == 0 else meta.add_row().cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)

    document.add_page_break()

    document.add_heading("Аннотация", level=1)
    document.add_paragraph(
        "В работе разработана локальная OSINT-система, которая собирает открытые данные Telegram-профилей, "
        "сохраняет их в SQLite, обогащает внешними аккаунтами через Sherlock/Snoop/Maigret и выполняет AI/NLP-анализ. "
        "Система включает desktop-интерфейс, веб-дашборд, экспорт отчетов, семантический поиск, классификацию "
        "профилей и модель проверки принадлежности найденных аккаунтов тому же человеку."
    )

    document.add_heading("Цель и задачи", level=1)
    document.add_paragraph(
        "Цель работы - создать прикладную систему для локального сбора, хранения, обогащения и анализа "
        "открытых данных о Telegram-профилях."
    )
    add_bullets(
        document,
        [
            "реализовать сбор открытых профилей Telegram через API;",
            "организовать локальное хранение данных в SQLite;",
            "интегрировать Sherlock, опционально Snoop и ручной Maigret для поиска внешних аккаунтов;",
            "рассчитать OSINT-показатели цифрового следа;",
            "добавить NLP-поиск, AI-классификацию и identity matching;",
            "подготовить интерфейс, выгрузки и сценарий демонстрации.",
        ],
    )

    document.add_heading("Архитектура", level=1)
    add_table(
        document,
        ["Слой", "Назначение", "Файлы"],
        [
            ["Сбор данных", "Telegram-группы и отдельные профили", "telegram_service.py, parser.py"],
            ["Хранение", "SQLite-схема, аналитика, CSV/Markdown-отчеты", "database.py, exp_csv.py"],
            ["OSINT", "Поиск внешних аккаунтов по username", "sherlock_integration.py, snoop_integration.py, maigret_integration.py"],
            ["AI/NLP", "Поиск, классификация, identity matching", "nlp_search_engine.py, ai_classifier.py, identity_matcher.py"],
        ],
    )

    document.add_heading("Реализованные функции", level=1)
    add_bullets(
        document,
        [
            "сбор открытых Telegram-профилей и realtime-проверка по username, ссылке или user id;",
            "сохранение профилей, групп, статусов проверок и внешних аккаунтов;",
            "ручной запуск Maigret для углубленной проверки username;",
            "расчет osint_score и exposure_level;",
            "семантический и гибридный поиск по профилям;",
            "классификация профилей: backend, frontend, designer, devops, analyst_security, other;",
            "обучение и отчетность ML-модулей;",
            "экспорт baseline/enriched CSV, OSINT CSV и Markdown-сводок.",
        ],
    )

    document.add_heading("Методика анализа", level=1)
    document.add_paragraph(
        "NLP-поиск использует sentence-transformers и fallback на лексический поиск при недоступности модели. "
        "Гибридное ранжирование объединяет семантическую близость, OSINT-балл и совпадения терминов. "
        "Identity matching оценивает сходство Telegram-профиля и найденного внешнего аккаунта по username, "
        "типу сайта, источнику, пересечению bio/URL и риску коллизии."
    )
    document.add_paragraph("Формула гибридного ранжирования:")
    document.add_paragraph(
        "hybrid_score = semantic_score * 0.70 + osint_boost * 0.20 + match_bonus * 0.10"
    )

    document.add_heading("Результаты", level=1)
    add_table(
        document,
        ["Показатель", "Значение"],
        [
            ["Всего профилей", "4959"],
            ["Профилей с bio", "1955"],
            ["Профилей с username", "3812"],
            ["Профилей с фото", "3195"],
        ],
    )
    add_table(
        document,
        ["Метрика identity matching", "Значение"],
        [
            ["Accuracy", "0.6264"],
            ["ROC-AUC", "0.7382"],
            ["F1 same_person", "0.5390"],
            ["F1 different_person", "0.6860"],
        ],
    )

    document.add_heading("Безопасность и ограничения", level=1)
    add_bullets(
        document,
        [
            "реальные Telegram API-ключи и session-файлы не хранятся в проекте;",
            "совпадение username не является доказательством личности;",
            "результаты Sherlock/Snoop/Maigret требуют ручной интерпретации;",
            "проект предназначен для учебного анализа открытых данных.",
        ],
    )

    document.add_heading("Проверка и демонстрация", level=1)
    document.add_paragraph(
        "Для проверки добавлены smoke-тесты ядра анализа: python -m unittest discover -s tests. "
        "Демонстрационный сценарий описан в DEMO_GUIDE.md. Основной запуск выполняется через app.py, "
        "дополнительный веб-дашборд - через streamlit run streamlit_app.py."
    )

    document.add_heading("Вывод", level=1)
    document.add_paragraph(
        "В результате создана локальная OSINT-система полного цикла: сбор открытых данных, хранение, "
        "обогащение, аналитика, AI/NLP-поиск, классификация и формирование отчетов. Практическая часть "
        "готова к демонстрации и защите при наличии локальных зависимостей и, для live-сбора, Telegram-ключей."
    )

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "OSINT Desktop coursework"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build()
